"""
Convert PRODUCT_DOCUMENTATION.md to a properly formatted DOCX file.
Removes markdown formatting and applies proper Word styles.
"""

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    import re
except ImportError:
    print("ERROR: python-docx not installed")
    print("Please run: pip install python-docx")
    exit(1)

def clean_markdown(text):
    """Remove markdown formatting symbols"""
    # Remove bold/italic markers
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)  # bold+italic
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)      # bold
    text = re.sub(r'\*(.+?)\*', r'\1', text)          # italic
    text = re.sub(r'__(.+?)__', r'\1', text)          # bold
    text = re.sub(r'_(.+?)_', r'\1', text)            # italic

    # Remove code markers
    text = re.sub(r'`([^`]+)`', r'\1', text)

    # Remove links but keep text
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)

    # Remove horizontal rules
    text = re.sub(r'^-{3,}$', '', text, flags=re.MULTILINE)

    return text

def add_heading(doc, text, level):
    """Add a heading with proper formatting"""
    text = clean_markdown(text)
    heading = doc.add_heading(text, level=level)
    if level == 1:
        heading.style.font.size = Pt(24)
        heading.style.font.color.rgb = RGBColor(0, 0, 0)
    elif level == 2:
        heading.style.font.size = Pt(18)
        heading.style.font.color.rgb = RGBColor(31, 78, 121)
    elif level == 3:
        heading.style.font.size = Pt(14)
        heading.style.font.color.rgb = RGBColor(31, 78, 121)
    return heading

def add_paragraph(doc, text, is_bullet=False, indent_level=0):
    """Add a paragraph with proper formatting"""
    text = clean_markdown(text.strip())
    if not text or text == '---':
        return

    para = doc.add_paragraph(text)

    if is_bullet:
        para.style = 'List Bullet' if indent_level == 0 else f'List Bullet {indent_level + 1}'
    else:
        para.style = 'Normal'
        para.paragraph_format.space_after = Pt(6)

    return para

def process_markdown_file(md_path, docx_path):
    """Convert markdown file to DOCX"""
    print(f"Reading {md_path}...")

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = Document()

    # Set default styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    print("Converting to DOCX...")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines
        if not line:
            i += 1
            continue

        # Headings
        if line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], 4)

        # Bullet lists
        elif line.startswith('- '):
            text = line[2:]
            add_paragraph(doc, text, is_bullet=True, indent_level=0)
        elif line.startswith('  - '):
            text = line[4:]
            add_paragraph(doc, text, is_bullet=True, indent_level=1)
        elif line.startswith('    - '):
            text = line[6:]
            add_paragraph(doc, text, is_bullet=True, indent_level=2)

        # Tables (convert to simple text)
        elif line.startswith('|'):
            # Skip table formatting, just extract content
            cells = [cell.strip() for cell in line.split('|')[1:-1]]
            if cells and not all(c.startswith('-') for c in cells):
                add_paragraph(doc, ' | '.join(cells))

        # Code blocks (skip)
        elif line.startswith('```'):
            # Skip code blocks
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                i += 1

        # Regular paragraphs
        else:
            add_paragraph(doc, line)

        i += 1

    print(f"Saving to {docx_path}...")
    doc.save(docx_path)
    print("Conversion complete!")

if __name__ == '__main__':
    import os

    # Get paths
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    md_path = os.path.join(project_dir, 'PRODUCT_DOCUMENTATION.md')
    docx_path = os.path.join(project_dir, 'PRODUCT_DOCUMENTATION.docx')

    if not os.path.exists(md_path):
        print(f"ERROR: {md_path} not found")
        exit(1)

    process_markdown_file(md_path, docx_path)
    print(f"\nDOCX file created: {docx_path}")
