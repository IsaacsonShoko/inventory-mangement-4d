"""
Convert client-facing markdown documents to professional DOCX format.
Removes emojis and formats for professional presentation.
"""

try:
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    import re
    import os
except ImportError:
    print("ERROR: python-docx not installed")
    print("Please run: pip install python-docx")
    exit(1)

def remove_emojis(text):
    """Remove all emojis from text"""
    emoji_pattern = re.compile(
        "["
        "\U0001F600-\U0001F64F"
        "\U0001F300-\U0001F5FF"
        "\U0001F680-\U0001F6FF"
        "\U0001F1E0-\U0001F1FF"
        "\U00002702-\U000027B0"
        "\U000024C2-\U0001F251"
        "\U0001F900-\U0001F9FF"
        "\U0001FA70-\U0001FAFF"
        "]+",
        flags=re.UNICODE
    )
    return emoji_pattern.sub('', text)

def clean_markdown(text):
    """Remove markdown formatting and emojis"""
    text = remove_emojis(text)
    text = re.sub(r'\*\*\*(.+?)\*\*\*', r'\1', text)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text)
    text = re.sub(r'^-{3,}$', '', text, flags=re.MULTILINE)
    text = re.sub(r'^_{3,}$', '', text, flags=re.MULTILINE)
    text = re.sub(r'\[ \]', '', text)
    text = re.sub(r'\[x\]', '', text, flags=re.IGNORECASE)
    return text.strip()

def add_heading(doc, text, level):
    """Add a heading with proper formatting"""
    text = clean_markdown(text)
    if not text:
        return
    heading = doc.add_heading(text, level=level)
    if level == 1:
        heading.style.font.size = Pt(24)
        heading.style.font.color.rgb = RGBColor(0, 0, 0)
        heading.style.font.bold = True
    elif level == 2:
        heading.style.font.size = Pt(18)
        heading.style.font.color.rgb = RGBColor(31, 78, 121)
        heading.style.font.bold = True
    elif level == 3:
        heading.style.font.size = Pt(14)
        heading.style.font.color.rgb = RGBColor(31, 78, 121)
    elif level == 4:
        heading.style.font.size = Pt(12)
        heading.style.font.color.rgb = RGBColor(68, 84, 106)
    return heading

def add_paragraph(doc, text, is_bullet=False, indent_level=0):
    """Add a paragraph with proper formatting"""
    text = clean_markdown(text)
    if not text or text == '---':
        return
    para = doc.add_paragraph(text)
    if is_bullet:
        para.style = 'List Bullet' if indent_level == 0 else f'List Bullet {indent_level + 1}'
    else:
        para.style = 'Normal'
        para.paragraph_format.space_after = Pt(6)
    return para

def process_table(doc, lines, start_index):
    """Process markdown table and add to document"""
    table_lines = []
    i = start_index
    while i < len(lines) and lines[i].strip().startswith('|'):
        line = lines[i].strip()
        if not re.match(r'^\|[\s\-\|:]+\|$', line):
            cells = [clean_markdown(cell.strip()) for cell in line.split('|')[1:-1]]
            if cells and any(cell for cell in cells):
                table_lines.append(cells)
        i += 1
    if table_lines and len(table_lines) > 0:
        num_cols = len(table_lines[0])
        table = doc.add_table(rows=len(table_lines), cols=num_cols)
        table.style = 'Light Grid Accent 1'
        for row_idx, row_data in enumerate(table_lines):
            for col_idx, cell_text in enumerate(row_data):
                if col_idx < len(table.rows[row_idx].cells):
                    table.rows[row_idx].cells[col_idx].text = cell_text
                    if row_idx == 0:
                        table.rows[row_idx].cells[col_idx].paragraphs[0].runs[0].font.bold = True
    return i - start_index

def process_markdown_file(md_path, docx_path):
    """Convert markdown file to DOCX"""
    print(f"Reading {os.path.basename(md_path)}...")
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    print("Converting to DOCX...")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith('# '):
            add_heading(doc, line[2:], 1)
        elif line.startswith('## '):
            add_heading(doc, line[3:], 2)
        elif line.startswith('### '):
            add_heading(doc, line[4:], 3)
        elif line.startswith('#### '):
            add_heading(doc, line[5:], 4)
        elif line.startswith('- '):
            text = line[2:]
            add_paragraph(doc, text, is_bullet=True, indent_level=0)
        elif line.startswith('  - '):
            text = line[4:]
            add_paragraph(doc, text, is_bullet=True, indent_level=1)
        elif line.startswith('    - '):
            text = line[6:]
            add_paragraph(doc, text, is_bullet=True, indent_level=2)
        elif line.startswith('|'):
            lines_processed = process_table(doc, lines, i)
            i += lines_processed
            continue
        elif line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                i += 1
        else:
            add_paragraph(doc, line)
        i += 1
    print(f"Saving to {os.path.basename(docx_path)}...")
    doc.save(docx_path)
    print(f"  Created: {docx_path}")

def main():
    """Convert all client documents"""
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    docs_dir = os.path.join(project_dir, 'docs')
    documents = [
        ('SYSTEM_OVERVIEW.md', 'SYSTEM_OVERVIEW.docx'),
    ]
    print("\n" + "="*60)
    print("XLINK CLIENT DOCUMENT CONVERTER")
    print("="*60 + "\n")
    converted = 0
    failed = 0
    for md_file, docx_file in documents:
        md_path = os.path.join(docs_dir, md_file)
        docx_path = os.path.join(docs_dir, docx_file)
        if not os.path.exists(md_path):
            print(f"SKIP: {md_file} (not found)")
            failed += 1
            continue
        try:
            process_markdown_file(md_path, docx_path)
            converted += 1
        except Exception as e:
            print(f"ERROR converting {md_file}: {e}")
            failed += 1
        print()
    print("="*60)
    print(f"CONVERSION COMPLETE")
    print(f"  Successful: {converted}")
    print(f"  Failed: {failed}")
    print("="*60)
    if converted > 0:
        print(f"\nDOCX files created in: {docs_dir}")

if __name__ == '__main__':
    main()
